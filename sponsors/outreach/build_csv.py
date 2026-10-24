"""Check the outreach drafts of one week and build the mail-merge CSV.

Usage (from the repo root):  python3 sponsors/outreach/build_csv.py week-1

Reads sponsors/outreach/<week>/NN-company.md (first line "Assunto:" or "Objet :"
style subject, then the body) and joins each draft to row NN of the
"This week's 40" tab and sponsors/pipeline.csv. Writes <week>/<week>.csv.
Exits 1 if any draft breaks the outreach rules (prompt 02, Part C), if a file
name does not match the company in row NN, or if a company is excluded (D9).
"""
import csv, pathlib, re, sys, unicodedata

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
MAX_WORDS = 119  # "under 120 words"
OPT_OUTS = (
    "Se preferir não receber mais emails nossos, basta responder a dizer.",
    "Si vous préférez ne plus recevoir nos messages, il suffit de répondre.",
    "Se preferir não receber mais nossos emails, é só responder.",
)
BANNED = ("paraíso", "paradis", "oportunidade única", "opportunité unique", "garantid", "garanti", "rentabilidade", "rendement")
# Notes for specific companies (by name: IDs change when the list is recomputed).
D6_NOTE = "Pipeline offer is '{offer}': pitch as a flat sponsorship only (D6); an affiliate deal needs a CEO decision."
EXTRA = {
    "Rio Credits": "Check the company is still active.",
}


def slugify(name: str) -> str:
    """File-name slug of a company: "Império (Ageas Portugal)" -> "imperio"."""
    name = re.sub(r"\(.*?\)", "", name)
    name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def load_week_rows() -> dict:
    """Rows of the "This week's 40" tab, keyed by their # column."""
    ws = openpyxl.load_workbook(ROOT / "sponsors" / "Sponsors_v2.xlsx", data_only=True)["This week's 40"]
    return {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[1]}


def main(week: str) -> int:
    folder = ROOT / "sponsors" / "outreach" / week
    rows = load_week_rows()
    pipe = {r["id"]: r for r in csv.DictReader(open(ROOT / "sponsors" / "pipeline.csv", encoding="utf-8"))}
    excl = openpyxl.load_workbook(ROOT / "sponsors" / "Sponsors_v2.xlsx", data_only=True)["Excluded (agencies)"]
    excluded = {slugify(r[0]) for r in excl.iter_rows(min_row=2, values_only=True) if r[0]}

    out, errors = [], []
    for f in sorted(folder.glob("[0-9][0-9]-*.md")):
        nn = int(f.name[:2])
        if nn not in rows:
            errors.append(f"{f.name}: no row {nn} in the tab")
            continue
        sid = rows[nn][1]
        p = pipe[sid]
        if f.stem[3:] != slugify(p["company"]):
            errors.append(f"{f.name}: row {nn} is {p['company']}, expected {nn:02d}-{slugify(p['company'])}.md")
        if slugify(p["company"]) in excluded:
            errors.append(f"{f.name}: {p['company']} is in the Excluded (agencies) tab (D9)")
        text = f.read_text(encoding="utf-8").strip()
        head, body = text.split("\n", 1)
        prefix, subject = (s.strip() for s in head.split(":", 1))
        body = body.strip()
        lang = "fr-FR" if prefix == "Objet" else "pt-BR" if p["edition"] == "BR" else "pt-PT"

        words = len(body.split())
        if words > MAX_WORDS:
            errors.append(f"{f.name}: {words} words (max {MAX_WORDS})")
        if re.search(r"^\s*([-*•]|\d+[.)])\s", body, re.M):
            errors.append(f"{f.name}: bullet or numbered list")
        if not any(o in body for o in OPT_OUTS):
            errors.append(f"{f.name}: opt-out line missing")
        if not body.endswith("[Assinatura]"):
            errors.append(f"{f.name}: must end with [Assinatura]")
        if body.count("?") != 1:
            errors.append(f"{f.name}: expected exactly one question (the call ask)")
        if "@" in text:
            errors.append(f"{f.name}: contains an email address")
        for b in BANNED:
            if b in text.lower():
                errors.append(f"{f.name}: banned word '{b}'")

        checks = []
        if p["evidence_year"] and int(p["evidence_year"]) < 2022:
            checks.append(f"Evidence is from {p['evidence_year']}: re-check it before sending (METHOD.md).")
        if p["evidence_status"] == "Search result":
            checks.append("Evidence page not read in full: open it before sending.")
        if "affiliate" in p["offer_type"].lower():
            checks.append(D6_NOTE.format(offer=p["offer_type"]))
        if p["company"] in EXTRA:
            checks.append(EXTRA[p["company"]])

        out.append({
            "nn": f"{nn:02d}",
            "id": sid,
            "company": p["company"],
            "edition": p["edition"],
            "language": lang,
            "recipient": f"{p['decision_maker_role']}. Find via: {p['contact_channel']}",
            "subject": subject,
            "body": body,
            "words": words,
            "send_on": rows[nn][10],
            "evidence_url": p["evidence_url"],
            "check_before_sending": " ".join(checks),
        })

    if len(out) != len(rows):
        errors.append(f"{len(out)} drafts for {len(rows)} companies in the tab")
    with open(folder / f"{week}.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(out[0].keys()))
        w.writeheader()
        w.writerows(out)
    print(f"{len(out)} drafts, {max(r['words'] for r in out)} words max -> {folder / (week + '.csv')}")
    for e in errors:
        print("ERROR", e)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "week-1"))

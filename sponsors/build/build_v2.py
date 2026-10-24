import csv, sys, datetime as dt
import os
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from sponsors_v2_data import ROWS, C_ROWS, PARTNERS, EXIT_BUYERS
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

REPO = os.path.dirname(HERE)

PRICES = {"Gold": 4800, "Silver": 1800, "Bronze": 540, "Campaign": 1500,
          "Featured listings": 1200, "Affiliate": 600, "Affiliate + sponsor": 1800}
PROB = {"A_recent": 0.25, "A_old": 0.15, "B": 0.12, "C": 0.04, "small_bonus": 0.05}
SPEED = {"S": 1.5, "M": 1.0, "L": 0.6}

def role(sector):
    s = sector.lower()
    if "bank" in s or "mortgage" in s: return "Marketing lead, non-resident / emigrant segment"
    if "credit" in s: return "Partnerships or marketing lead"
    if "fx" in s or "fintech" in s or "nif" in s: return "Partnerships / affiliates manager"
    if "law" in s or "legal" in s or "residency" in s or "visa" in s or "immigration" in s or "succession" in s: return "Managing partner or business development"
    if "insurance" in s: return "Marketing lead, retail"
    if "developer" in s or "builder" in s: return "Sales or marketing director"
    if "agency" in s or "broker" in s or "portal" in s or "consultancy" in s or "buying" in s: return "Owner / agency director"
    return "Owner or marketing lead"

def channel(size, tier):
    if size == "S": return "Phone or email from company website, then LinkedIn"
    if size == "M": return "LinkedIn (named role), then company email"
    return "LinkedIn (named role); general partnerships email as backup"

# ---------- assemble ----------
recs = []
for r in ROWS:
    company, sector, edition, tier, url, note, year, status, size, offer = r
    recs.append(dict(company=company, sector=sector, edition=edition, tier=tier, size=size,
                     evidence_url=url, evidence_note=note, evidence_year=year, evidence_status=status,
                     offer_type=offer))
for company, sector in C_ROWS:
    recs.append(dict(company=company, sector=sector, edition="LU, FR, CH, BR" if "(" not in sector else
                     ("LU" if "Luxembourg" in sector else "CH" if "Switzerland" in sector else "BR" if "Brazil" in sector else "FR"),
                     tier="C", size="L", evidence_url="", evidence_note="Large or generic; no proof of spend on our audience",
                     evidence_year="", evidence_status="", offer_type="Gold" if "Bank" in sector or "Airline" in sector else "Silver"))

names = [r["company"] for r in recs]
assert len(names) == len(set(names)), "duplicate company"

def calc(r):
    v = PRICES[r["offer_type"]]
    if r["tier"] == "A": p = PROB["A_recent"] if int(r["evidence_year"]) >= 2022 else PROB["A_old"]
    elif r["tier"] == "B": p = PROB["B"]
    else: p = PROB["C"]
    if r["size"] == "S": p += PROB["small_bonus"]
    ev = v * p
    return v, round(p, 2), round(ev, 2), round(ev * SPEED[r["size"]], 2)

for r in recs:
    r["est_deal_value_eur"], r["close_probability"], r["expected_value_eur"], r["priority_score"] = calc(r)
    r["decision_maker_role"] = role(r["sector"])
    r["contact_channel"] = channel(r["size"], r["tier"])

ab = [r for r in recs if r["tier"] in "AB"]
cc = [r for r in recs if r["tier"] == "C"]
ab.sort(key=lambda r: (-r["priority_score"], r["company"]))

# schedule: top 40 over 26-30 Oct, next 50 over 2-6 Nov
days1 = [dt.date(2026, 10, 26) + dt.timedelta(d) for d in range(5)]
days2 = [dt.date(2026, 11, 2) + dt.timedelta(d) for d in range(5)]
for i, r in enumerate(ab):
    r["status"] = "To contact"
    r["owner"] = "Partner"
    if i < 40:
        r["next_action_date"] = days1[i // 8].isoformat()
        r["next_action"] = "Email 1 (template in sponsors/METHOD.md) + LinkedIn request on day 2"
    else:
        r["next_action_date"] = days2[(i - 40) // 10].isoformat()
        r["next_action"] = "Email 1 + LinkedIn request on day 2"
for r in cc:
    r["status"] = "Parked"; r["owner"] = "CEO"
    r["next_action"] = "Revisit when 1,000 subscribers or a warm intro exists"
    r["next_action_date"] = "2026-12-01"

ordered = ab + sorted(cc, key=lambda r: r["company"])
for i, r in enumerate(ordered, 1):
    r["id"] = f"S{i:03d}"

COLS = ["id", "company", "sector", "edition", "tier", "size", "evidence_url", "evidence_note", "evidence_year",
        "evidence_status", "decision_maker_role", "contact_channel", "offer_type", "est_deal_value_eur",
        "close_probability", "expected_value_eur", "priority_score", "status", "next_action", "next_action_date", "owner"]

with open(f"{REPO}/pipeline.csv", "w", newline="", encoding="utf-8") as fh:
    w = csv.DictWriter(fh, fieldnames=COLS, extrasaction="ignore")
    w.writeheader()
    for r in ordered: w.writerow(r)

# ---------- workbook ----------
AR = "Arial"
H = Font(name=AR, bold=True, color="FFFFFF", size=10)
HF = PatternFill("solid", fgColor="16202C")
B = Font(name=AR, size=10); BB = Font(name=AR, size=10, bold=True)
BLUE = Font(name=AR, size=10, color="0000FF"); GREEN = Font(name=AR, size=10, color="008000")
YEL = PatternFill("solid", fgColor="FFFF00")
TA = PatternFill("solid", fgColor="D6E4EF"); TB = PatternFill("solid", fgColor="EEF3F7")
thin = Side(style="thin", color="DFE3E8"); BD = Border(bottom=thin)
WR = Alignment(wrap_text=True, vertical="top")
EUR = '#,##0 "€";(#,##0 "€");"-"'
PCT = '0%;(0%);"-"'

wb = Workbook()

# Settings
st = wb.active; st.title = "Settings"
st["A1"] = "Settings: edit the yellow cells; every value in the workbook updates"; st["A1"].font = Font(name=AR, bold=True, size=12)
st["A3"], st["B3"], st["C3"] = "Offer type", "Price year 1 (€)", "Basis"
for c in ("A3", "B3", "C3"): st[c].font = H; st[c].fill = HF
basis = {"Gold": "8,000 €/yr list price, 40% founding discount", "Silver": "3,000 €/yr list, 40% founding discount",
         "Bronze": "900 €/yr list, 40% founding discount", "Campaign": "One developer campaign",
         "Featured listings": "~5 listings x 12 months at reduced pilot rate", "Affiliate": "Estimated affiliate income per year",
         "Affiliate + sponsor": "Affiliate income plus a Silver slot"}
for i, (k, v) in enumerate(PRICES.items(), 4):
    st.cell(row=i, column=1, value=k).font = B
    c = st.cell(row=i, column=2, value=v); c.font = BLUE; c.fill = YEL; c.number_format = EUR
    st.cell(row=i, column=3, value=basis[k]).font = B
pr_end = 3 + len(PRICES)
r0 = pr_end + 2
st.cell(row=r0, column=1, value="Close probability").font = H; st.cell(row=r0, column=1).fill = HF
st.cell(row=r0, column=2, value="Value").font = H; st.cell(row=r0, column=2).fill = HF
st.cell(row=r0, column=3, value="Meaning").font = H; st.cell(row=r0, column=3).fill = HF
prob_rows = [("Tier A, evidence 2022 or later", PROB["A_recent"], "Proven recent spend on our audience"),
             ("Tier A, older evidence", PROB["A_old"], "Proven spend before 2022; re-verify"),
             ("Tier B", PROB["B"], "Serves our audience, no proof of spend"),
             ("Tier C", PROB["C"], "Large or generic, long cycle"),
             ("Bonus for small company", PROB["small_bonus"], "Owner decides fast")]
P = {}
for i, (k, v, m) in enumerate(prob_rows, r0 + 1):
    st.cell(row=i, column=1, value=k).font = B
    c = st.cell(row=i, column=2, value=v); c.font = BLUE; c.fill = YEL; c.number_format = PCT
    st.cell(row=i, column=3, value=m).font = B
    P[k] = f"Settings!$B${i}"
s0 = r0 + len(prob_rows) + 2
st.cell(row=s0, column=1, value="Size").font = H; st.cell(row=s0, column=1).fill = HF
st.cell(row=s0, column=2, value="Speed factor").font = H; st.cell(row=s0, column=2).fill = HF
st.cell(row=s0, column=3, value="Meaning").font = H; st.cell(row=s0, column=3).fill = HF
for i, (k, v, m) in enumerate([("S", 1.5, "Small: decides in days"), ("M", 1.0, "Mid-size: weeks"), ("L", 0.6, "Large: months")], s0 + 1):
    st.cell(row=i, column=1, value=k).font = B
    c = st.cell(row=i, column=2, value=v); c.font = BLUE; c.fill = YEL; c.number_format = "0.0x"
    st.cell(row=i, column=3, value=m).font = B
sp_start, sp_end = s0 + 1, s0 + 3
n0 = sp_end + 2
st.cell(row=n0, column=1, value="Prices follow the current plan. If Daniel picks lower pre-launch pricing (decision D3), change column B above.").font = Font(name=AR, size=9, italic=True, color="5D6672")
st.column_dimensions["A"].width = 34; st.column_dimensions["B"].width = 18; st.column_dimensions["C"].width = 48

# Pipeline sheet (A/B), Later sheet (C)
HEAD = ["ID", "Company", "Sector", "Editions", "Tier", "Size", "Evidence URL", "Evidence", "Evidence year", "Evidence check",
        "Who to reach", "Channel", "Offer", "Deal value year 1 (€)", "Close probability", "Expected value (€)", "Priority score",
        "Status", "Next action", "Next action date", "Owner"]
WID = [7, 30, 24, 14, 6, 6, 34, 50, 9, 13, 30, 30, 16, 13, 11, 13, 11, 14, 36, 13, 10]
STATUSES = '"To contact,Contacted,Replied,Call booked,Proposal sent,Won,Lost,Parked"'

def sheet(name, rows):
    ws = wb.create_sheet(name)
    for i, (h, w) in enumerate(zip(HEAD, WID), 1):
        c = ws.cell(row=1, column=i, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"
    for n, r in enumerate(rows, 2):
        vals = [r["id"], r["company"], r["sector"], r["edition"], r["tier"], r["size"], r["evidence_url"] or None,
                r["evidence_note"], r["evidence_year"] or None, r["evidence_status"] or None, r["decision_maker_role"],
                r["contact_channel"], r["offer_type"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=n, column=c, value=v); cell.font = BB if c == 2 else B; cell.alignment = WR; cell.border = BD
        if r["evidence_url"]:
            ws.cell(row=n, column=7).hyperlink = r["evidence_url"]; ws.cell(row=n, column=7).font = Font(name=AR, size=10, color="1F5F8B", underline="single")
        ws.cell(row=n, column=5).fill = TA if r["tier"] == "A" else TB if r["tier"] == "B" else PatternFill()
        # formulas
        ws.cell(row=n, column=14, value=f"=INDEX(Settings!$B$4:$B${pr_end},MATCH(M{n},Settings!$A$4:$A${pr_end},0))").number_format = EUR
        ws.cell(row=n, column=15, value=(f'=IF(E{n}="A",IF(I{n}>=2022,{P["Tier A, evidence 2022 or later"]},{P["Tier A, older evidence"]}),'
                                          f'IF(E{n}="B",{P["Tier B"]},{P["Tier C"]}))+IF(F{n}="S",{P["Bonus for small company"]},0)')).number_format = PCT
        ws.cell(row=n, column=16, value=f"=N{n}*O{n}").number_format = EUR
        ws.cell(row=n, column=17, value=f"=P{n}*INDEX(Settings!$B${sp_start}:$B${sp_end},MATCH(F{n},Settings!$A${sp_start}:$A${sp_end},0))").number_format = '#,##0'
        for c in (14, 15, 16, 17):
            ws.cell(row=n, column=c).font = B; ws.cell(row=n, column=c).border = BD; ws.cell(row=n, column=c).alignment = Alignment(vertical="top")
        for c, v in ((18, r["status"]), (19, r["next_action"]), (20, r["next_action_date"]), (21, r["owner"])):
            cell = ws.cell(row=n, column=c, value=v); cell.font = B; cell.alignment = WR; cell.border = BD
    last = len(rows) + 1
    dv = DataValidation(type="list", formula1=STATUSES, allow_blank=False)
    ws.add_data_validation(dv); dv.add(f"R2:R{last}")
    ws.auto_filter.ref = f"A1:U{last}"
    return ws, last

pl, pl_last = sheet("Pipeline A-B", ab)
lt, lt_last = sheet("Later (tier C)", sorted(cc, key=lambda r: r["company"]))

# This week's 40
tw = wb.create_sheet("This week's 40")
th = ["#", "ID", "Company", "Editions", "Tier", "Offer", "Expected value (€)", "Priority score", "Who to reach", "Channel", "Send on", "Status"]
tws = [5, 7, 32, 14, 6, 16, 14, 12, 34, 34, 12, 14]
for i, (h, w) in enumerate(zip(th, tws), 1):
    c = tw.cell(row=1, column=i, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
    tw.column_dimensions[get_column_letter(i)].width = w
tw.freeze_panes = "D2"
rng = f"'Pipeline A-B'!$A$2:$A${pl_last}"
def look(col, n): return f"=INDEX('Pipeline A-B'!${col}$2:${col}${pl_last},MATCH($B{n},{rng},0))"
for k, r in enumerate(ab[:40], 2):
    tw.cell(row=k, column=1, value=k - 1).font = B
    tw.cell(row=k, column=2, value=r["id"]).font = B
    for col, src in ((3, "B"), (4, "D"), (5, "E"), (6, "M"), (7, "P"), (8, "Q"), (9, "K"), (10, "L"), (11, "T"), (12, "R")):
        c = tw.cell(row=k, column=col, value=look(src, k)); c.font = BB if col == 3 else B; c.alignment = WR; c.border = BD
    tw.cell(row=k, column=7).number_format = EUR; tw.cell(row=k, column=8).number_format = '#,##0'
n = 43
tw.cell(row=n, column=3, value="Order is a snapshot of priority on 24 Oct 2026. Details update live from 'Pipeline A-B'; to re-rank, sort that sheet by Priority score.").font = Font(name=AR, size=9, italic=True, color="5D6672")

# Partners & media
pm = wb.create_sheet("Partners & media")
ph = ["Organisation", "Edition", "Why it matters", "Evidence URL", "Next action"]
for i, (h, w) in enumerate(zip(ph, [42, 10, 60, 40, 40]), 1):
    c = pm.cell(row=1, column=i, value=h); c.font = H; c.fill = HF
    pm.column_dimensions[get_column_letter(i)].width = w
for k, p in enumerate(PARTNERS, 2):
    for c, v in enumerate(p, 1):
        cell = pm.cell(row=k, column=c, value=v); cell.font = BB if c == 1 else B; cell.alignment = WR; cell.border = BD
    if p[3]: pm.cell(row=k, column=4).hyperlink = p[3]
pm.cell(row=len(PARTNERS) + 3, column=1, value="Not sponsors: they give access to sponsors, audience or intelligence on who buys diaspora ads.").font = Font(name=AR, size=9, italic=True, color="5D6672")

# Exit buyers
eb = wb.create_sheet("Exit buyers")
eh = ["Possible buyer", "Type", "Why they would buy", "Evidence URL"]
for i, (h, w) in enumerate(zip(eh, [32, 26, 66, 50]), 1):
    c = eb.cell(row=1, column=i, value=h); c.font = H; c.fill = HF
    eb.column_dimensions[get_column_letter(i)].width = w
for k, p in enumerate(EXIT_BUYERS, 2):
    for c, v in enumerate(p, 1):
        cell = eb.cell(row=k, column=c, value=v or None); cell.font = BB if c == 1 else B; cell.alignment = WR; cell.border = BD
    if p[3]: eb.cell(row=k, column=4).hyperlink = p[3]

# Summary
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "Sponsor pipeline v2: follow the money"; sm["A1"].font = Font(name=AR, bold=True, size=13)
sm["A2"] = "Built 24 Oct 2026. Tier A = proven spend on our audience; B = serves our audience; C = large/generic (parked)."; sm["A2"].font = Font(name=AR, size=9, italic=True, color="5D6672")
hdr = ["", "Companies", "Expected value (€)", "Won (€)"]
for i, h in enumerate(hdr, 1):
    c = sm.cell(row=4, column=i, value=h); c.font = H; c.fill = HF
PLR = f"'Pipeline A-B'"; LTR = f"'Later (tier C)'"
rows_sum = [("Tier A", "A"), ("Tier B", "B")]
r = 5
for label, t in rows_sum:
    sm.cell(row=r, column=1, value=label).font = B
    sm.cell(row=r, column=2, value=f'=COUNTIF({PLR}!$E$2:$E${pl_last},"{t}")').font = B
    c = sm.cell(row=r, column=3, value=f'=SUMIF({PLR}!$E$2:$E${pl_last},"{t}",{PLR}!$P$2:$P${pl_last})'); c.font = B; c.number_format = EUR
    c = sm.cell(row=r, column=4, value=f'=SUMIFS({PLR}!$N$2:$N${pl_last},{PLR}!$E$2:$E${pl_last},"{t}",{PLR}!$R$2:$R${pl_last},"Won")'); c.font = B; c.number_format = EUR
    r += 1
sm.cell(row=r, column=1, value="Tier C (parked)").font = B
sm.cell(row=r, column=2, value=f"=COUNTA({LTR}!$A$2:$A${lt_last})").font = B
c = sm.cell(row=r, column=3, value=f"=SUM({LTR}!$P$2:$P${lt_last})"); c.font = B; c.number_format = EUR
c = sm.cell(row=r, column=4, value=f'=SUMIF({LTR}!$R$2:$R${lt_last},"Won",{LTR}!$N$2:$N${lt_last})'); c.font = B; c.number_format = EUR
r += 1
sm.cell(row=r, column=1, value="Total").font = BB
for col in (2, 3, 4):
    L = get_column_letter(col)
    c = sm.cell(row=r, column=col, value=f"=SUM({L}5:{L}{r-1})"); c.font = BB
    if col > 2: c.number_format = EUR
tot = r
r += 1
sm.cell(row=r, column=1, value="Share of tier A + B").font = B
c = sm.cell(row=r, column=2, value=f"=(B5+B6)/B{tot}"); c.font = B; c.number_format = "0%"

r += 2
for i, h in enumerate(["Edition", "A + B companies", "Expected value (€)"], 1):
    c = sm.cell(row=r, column=i, value=h); c.font = H; c.fill = HF
for ed in ["LU", "FR", "CH", "BR"]:
    r += 1
    sm.cell(row=r, column=1, value=ed).font = B
    sm.cell(row=r, column=2, value=f'=SUMPRODUCT(--ISNUMBER(SEARCH("{ed}",{PLR}!$D$2:$D${pl_last})))').font = B
    c = sm.cell(row=r, column=3, value=f'=SUMPRODUCT(ISNUMBER(SEARCH("{ed}",{PLR}!$D$2:$D${pl_last}))*{PLR}!$P$2:$P${pl_last})'); c.font = B; c.number_format = EUR
r += 1
sm.cell(row=r, column=1, value="A company with several editions counts in each, so edition rows add up to more than the total.").font = Font(name=AR, size=9, italic=True, color="5D6672")

r += 2
for i, h in enumerate(["Status", "A + B companies"], 1):
    c = sm.cell(row=r, column=i, value=h); c.font = H; c.fill = HF
for s in ["To contact", "Contacted", "Replied", "Call booked", "Proposal sent", "Won", "Lost"]:
    r += 1
    sm.cell(row=r, column=1, value=s).font = B
    sm.cell(row=r, column=2, value=f'=COUNTIF({PLR}!$R$2:$R${pl_last},"{s}")').font = B
sm.column_dimensions["A"].width = 30
for L in "BCD": sm.column_dimensions[L].width = 20

wb.move_sheet("Settings", offset=len(wb.sheetnames))
wb.active = 0
wb.save(f"{REPO}/Sponsors_v2.xlsx")
print("rows", len(ordered), "A", sum(r['tier']=='A' for r in ordered), "B", sum(r['tier']=='B' for r in ordered), "C", len(cc))
print("top10:"); [print(" ", r["id"], r["company"], r["tier"], r["size"], r["offer_type"], r["priority_score"]) for r in ab[:10]]
